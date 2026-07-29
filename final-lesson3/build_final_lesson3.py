from __future__ import annotations

import argparse
import copy
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from lesson3_data import (
    ANALYSIS_NAMES,
    ANALYSIS_Q,
    AUTHOR,
    COMPLETE_MEANING,
    COMPLETE_TEXT,
    GENERAL_IDEA,
    GENERAL_QA,
    LESSON_TITLE,
    MCQ,
    PARAGRAPHS,
    SCHOOLBOOK_END,
    TASTE_NAMES,
    TASTE_Q,
)

FONT = "Arial"
DARK_BLUE = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
ORANGE = "C65911"
LIGHT_ORANGE = "FCE4D6"
GREEN = "548235"
LIGHT_GREEN = "E2F0D9"
RED = "C00000"
LIGHT_RED = "FDE9E7"
PURPLE = "7030A0"
LIGHT_PURPLE = "E4DFEC"
GRAY = "666666"
LIGHT_GRAY = "F2F2F2"
GOLD = "BF9000"
WHITE = "FFFFFF"
BLACK = "000000"


def set_bidi_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = align
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.line_spacing = 1.18


def set_run(run, size=14, bold=False, color=BLACK, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:cs"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().append(OxmlElement("w:rtl"))
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="A6A6A6", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_rtl(table):
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    table.autofit = True
    tbl_pr = table._tbl.tblPr
    bidi = tbl_pr.find(qn("w:bidiVisual"))
    if bidi is None:
        bidi = OxmlElement("w:bidiVisual")
        tbl_pr.append(bidi)
    bidi.set(qn("w:val"), "1")
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        bidi_row = tr_pr.find(qn("w:bidiVisual"))
        if bidi_row is None:
            bidi_row = OxmlElement("w:bidiVisual")
            tr_pr.append(bidi_row)
        bidi_row.set(qn("w:val"), "1")
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_bidi = tc_pr.find(qn("w:bidi"))
            if tc_bidi is None:
                tc_bidi = OxmlElement("w:bidi")
                tc_pr.append(tc_bidi)
            tc_bidi.set(qn("w:val"), "1")
            for p in cell.paragraphs:
                set_bidi_paragraph(p)


def clear_cell(cell):
    cell.text = ""
    p = cell.paragraphs[0]
    set_bidi_paragraph(p)
    return p


def add_cell_text(cell, text, size=12.5, bold=False, color=BLACK, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p = clear_cell(cell)
    p.alignment = align
    set_bidi_paragraph(p, align)
    set_run(p.add_run(str(text)), size=size, bold=bold, color=color)
    return p


def add_title_band(doc, text, fill=DARK_BLUE, size=20, page_break_before=False):
    if page_break_before:
        doc.add_page_break()
    table = doc.add_table(rows=1, cols=1)
    set_table_rtl(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, fill)
    add_cell_text(cell, text, size=size, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()


def add_subtitle(doc, text, color=BLUE, size=16):
    p = doc.add_paragraph()
    set_bidi_paragraph(p)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run(text), size=size, bold=True, color=color)
    return p


def add_box(doc, text, fill=LIGHT_BLUE, border=BLUE, label=None, size=13.5):
    table = doc.add_table(rows=1, cols=1)
    set_table_rtl(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, border, "12")
    p = clear_cell(cell)
    if label:
        set_run(p.add_run(label + " "), size=size, bold=True, color=border)
    set_run(p.add_run(text), size=size, color=BLACK)
    return table


def add_body_paragraph(doc, text, bullet=False, size=13.3, color=BLACK, bold=False):
    p = doc.add_paragraph()
    set_bidi_paragraph(p)
    p.paragraph_format.first_line_indent = Cm(0.5) if not bullet else Cm(0)
    p.paragraph_format.left_indent = Cm(0.15)
    p.paragraph_format.right_indent = Cm(0.15)
    prefix = "• " if bullet else ""
    set_run(p.add_run(prefix + text), size=size, color=color, bold=bold)
    return p


def add_table(doc, headers, rows, header_fill=DARK_BLUE, alt_fill=LIGHT_GRAY, font_size=11.7):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_rtl(table)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, header_fill)
        add_cell_text(cell, header, size=font_size, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for c_index, value in enumerate(row_data):
            if r_index % 2:
                set_cell_shading(cells[c_index], alt_fill)
            add_cell_text(cells[c_index], value, size=font_size)
    set_table_rtl(table)
    doc.add_paragraph()
    return table


def add_question(doc, number, question, answer=None, source=None, answer_mode=False, lines=2):
    p = doc.add_paragraph()
    set_bidi_paragraph(p)
    set_run(p.add_run(f"{number} ـ {question}"), size=13.2, bold=True, color=DARK_BLUE)
    if source:
        set_run(p.add_run(f"  المصدر: {source}"), size=10.5, bold=True, color=RED)
    if answer_mode and answer is not None:
        ap = doc.add_paragraph()
        set_bidi_paragraph(ap)
        set_run(ap.add_run("الإجابة: "), size=12.7, bold=True, color=GREEN)
        set_run(ap.add_run(answer), size=12.7, color=BLACK)
    else:
        for _ in range(lines):
            lp = doc.add_paragraph("............................................................................................................................")
            set_bidi_paragraph(lp)
            for run in lp.runs:
                set_run(run, size=11, color=GRAY)


def add_qa_section(doc, title, qa, answer_mode=False, source=None):
    add_subtitle(doc, title, color=BLUE)
    for i, item in enumerate(qa, 1):
        q, a = item
        add_question(doc, str(i), q, a, source=source, answer_mode=answer_mode)


def paragraph_exercise_items(pdata, index):
    word0 = pdata["vocab"][0]
    word1 = pdata["vocab"][1]
    first_expr = pdata["expressions"][0]
    first_value = pdata["values"][0]
    analysis = pdata["analysis"]
    taste = pdata["taste"]
    book_q, book_a = pdata["book"][0]
    return [
        (f"هات معنى {word0[0]} ومضاد {word1[0]}.", f"معنى {word0[0]}: {word0[1]}. مضاد {word1[0]}: {word1[2]}."),
        ("صغ الفكرة الجزئية للفقرة بأسلوبك.", pdata["idea"]),
        ("اشرح الفقرة شرحًا مركزًا يوضح رسالتها.", " ".join(pdata["explanation"][:2])),
        (book_q, book_a),
        ("حدد نوع النص في الفقرة واذكر سمة ظهرت فيها.", analysis[0][0] + " ومن سماته التوجيه والتفسير وترابط الأفكار."),
        ("حدد الفكرة الرئيسة ومعنى مباشرًا وآخر ضمنيًا.", analysis[1][0]),
        ("وضح بنية الفقرة وترتيب أفكارها.", analysis[2][0]),
        (f"وضح دلالة التعبير: {first_expr[0]}.", first_expr[1]),
        ("استخرج علاقة بين الكلمات أو الجمل، ووضح أثرها.", taste[2][0] + " " + taste[3][0]),
        ("حدد أسلوبًا واذكر غرضه.", taste[4][0]),
        ("احكم على الفقرة من حيث وضوحها وإقناعها.", analysis[5][0]),
        (f"اقترح تطبيقًا حياتيًا لقيمة {first_value[0]}.", first_value[3]),
    ]


def add_cover(doc, answer_mode):
    section = doc.sections[0]
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.3)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.header_distance = Cm(0.55)
    section.footer_distance = Cm(0.55)

    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER)
    set_run(p.add_run("اللغة العربية"), size=22, bold=True, color=ORANGE)
    p = doc.add_paragraph()
    set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER)
    set_run(p.add_run("الصف الثالث الإعدادي"), size=18, bold=True, color=BLUE)

    title_table = doc.add_table(rows=1, cols=1)
    set_table_rtl(title_table)
    cell = title_table.cell(0, 0)
    set_cell_shading(cell, DARK_BLUE)
    set_cell_border(cell, ORANGE, "18")
    add_cell_text(cell, LESSON_TITLE, size=28, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)

    p = doc.add_paragraph()
    set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER)
    set_run(p.add_run(AUTHOR), size=17, bold=True, color=GREEN)
    p = doc.add_paragraph()
    set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER)
    version = "نموذج الإجابة الكامل" if answer_mode else "نسخة الطالب"
    set_run(p.add_run(version), size=20, bold=True, color=RED if answer_mode else GOLD)

    add_box(doc, "شرح تفصيلي ـ قاموس شامل ـ تعبيرات ودلالات ـ قيم وتطبيقات ـ فهم واستيعاب ـ تحليل النص بعناصره ـ تذوق النص بعناصره ـ أسئلة كتاب المدرسة ـ تدريبات شاملة", fill=LIGHT_ORANGE, border=ORANGE, size=14)
    if not answer_mode:
        add_body_paragraph(doc, "اسم الطالب: ................................................................................................................", size=13)
        add_body_paragraph(doc, "الفصل: ........................................     التاريخ: ........................................", size=13)
    doc.add_page_break()


def add_header_footer(doc, answer_mode):
    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER)
        set_run(p.add_run(f"{LESSON_TITLE} ـ {'نموذج الإجابة' if answer_mode else 'نسخة الطالب'}"), size=9.5, bold=True, color=GRAY)
        footer = section.footer
        fp = footer.paragraphs[0]
        set_bidi_paragraph(fp, WD_ALIGN_PARAGRAPH.CENTER)
        set_run(fp.add_run("الصف الثالث الإعدادي"), size=9.5, color=GRAY)
        set_run(fp.add_run("   ـ   "), size=9.5, color=GRAY)
        fld_char1 = OxmlElement("w:fldChar")
        fld_char1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        fld_char2 = OxmlElement("w:fldChar")
        fld_char2.set(qn("w:fldCharType"), "end")
        r = fp.add_run()._r
        r.append(fld_char1)
        r.append(instr)
        r.append(fld_char2)


def add_intro(doc):
    add_title_band(doc, "مدخل إلى الدرس")
    add_box(doc, GENERAL_IDEA, fill=LIGHT_BLUE, border=BLUE, label="الفكرة العامة:")
    rows = [
        ["نوع النص", "نص فكري اجتماعي إرشادي؛ يعرض صفات الصديق وحقوقه ويدعو إلى حسن الاختيار."],
        ["المجال", "القيم الإنسانية والعلاقات الاجتماعية وبناء الشخصية."],
        ["هدف الكاتب", "توضيح معنى الصداقة الحقيقية، وتقويم السلوك، والتحذير من الصداقة الزائفة."],
        ["الجمهور المستهدف", "الشباب والطلاب والأبناء وكل من يكون علاقات وصداقات."],
        ["بنية النص", "ست فقرات: صفات وحقوق وآداب وتحذير وخاتمة تربط الصحبة بالمستقبل."],
        ["المغزى", "اختر صديقك بعناية، وكن أنت الصديق الذي تبحث عنه: أمينًا وفيًا رفيقًا مخلصًا."],
    ]
    add_table(doc, ["العنصر", "البيان"], rows, header_fill=BLUE, alt_fill=LIGHT_BLUE, font_size=12.3)
    add_subtitle(doc, "الأهداف التعليمية", color=ORANGE)
    goals = [
        "فهم أفكار النص وشرحها وربطها بحياة الطالب.",
        "تفسير المفردات وتحديد المعنى والمضاد والجمع والمفرد.",
        "تحليل كل فقرة وفق عناصر تحليل النص الثمانية.",
        "تذوق كل فقرة وفق عناصر تذوق النص الثلاثة عشر مع الشاهد والأثر.",
        "التمييز بين الحوار والمراء، والوفاء والمصلحة، والصداقة الحقيقية والزائفة.",
        "الإجابة عن أسئلة كتاب المدرسة والتدريبات المتنوعة والشاملة.",
    ]
    for g in goals:
        add_body_paragraph(doc, g, bullet=True)


def add_paragraph_unit(doc, pdata, idx, answer_mode):
    add_title_band(doc, pdata["title"], fill=DARK_BLUE, page_break_before=True)
    add_box(doc, pdata["text"], fill="FFFBE6", border=GOLD, label="النص:", size=14)
    add_box(doc, pdata["idea"], fill=LIGHT_BLUE, border=BLUE, label="الفكرة الجزئية:")

    add_subtitle(doc, "الشرح التفصيلي", color=ORANGE)
    for point in pdata["explanation"]:
        add_body_paragraph(doc, point, bullet=True)

    add_subtitle(doc, "القاموس اللغوي الشامل", color=BLUE)
    add_table(doc, ["الكلمة أو التعبير", "المعنى", "المضاد", "الجمع", "المفرد"], pdata["vocab"], header_fill=BLUE, alt_fill=LIGHT_BLUE, font_size=11.2)

    add_subtitle(doc, "التعبيرات ودلالاتها", color=RED)
    add_table(doc, ["التعبير", "الدلالة والأثر"], pdata["expressions"], header_fill=RED, alt_fill=LIGHT_RED, font_size=11.7)

    add_subtitle(doc, "القيم والمبادئ والتطبيقات الحياتية", color=GREEN)
    add_table(doc, ["القيمة", "المقصود بها", "الدليل من النص", "التطبيق الحياتي"], pdata["values"], header_fill=GREEN, alt_fill=LIGHT_GREEN, font_size=11.2)

    add_qa_section(doc, "أسئلة الفهم والاستيعاب", pdata["qa"], answer_mode=answer_mode)
    add_qa_section(doc, "أسئلة كتاب المدرسة المدمجة", pdata["book"], answer_mode=answer_mode, source="كتاب المدرسة")

    add_subtitle(doc, "أولًا ـ عناصر تحليل النص مطبقة على الفقرة", color=PURPLE, size=17)
    analysis_rows = []
    for name, content in zip(ANALYSIS_NAMES, pdata["analysis"]):
        analysis_rows.append([name, content[0]])
    add_table(doc, ["العنصر الصحيح", "التطبيق التفصيلي على الفقرة"], analysis_rows, header_fill=PURPLE, alt_fill=LIGHT_PURPLE, font_size=11.3)

    add_subtitle(doc, "ثانيًا ـ عناصر تذوق النص مطبقة على الفقرة", color=ORANGE, size=17)
    taste_rows = []
    for name, content in zip(TASTE_NAMES, pdata["taste"]):
        taste_rows.append([name, content[0]])
    add_table(doc, ["العنصر الصحيح", "الشاهد والتفسير والأثر"], taste_rows, header_fill=ORANGE, alt_fill=LIGHT_ORANGE, font_size=11.1)

    add_subtitle(doc, f"تدريب شامل على الفقرة {idx}", color=DARK_BLUE, size=17)
    exercise_items = paragraph_exercise_items(pdata, idx)
    for qnumb, (q, a) in enumerate(exercise_items, 1):
        source = "كتاب المدرسة" if qnumb == 4 else None
        add_question(doc, str(qnumb), q, a, source=source, answer_mode=answer_mode, lines=2)


def add_numbered_exercise(doc, title, items, answer_mode, kind="qa", page_break=True):
    add_title_band(doc, title, fill=DARK_BLUE, page_break_before=page_break)
    if kind == "mcq":
        for i, (stem, choices, answer) in enumerate(items, 1):
            p = doc.add_paragraph()
            set_bidi_paragraph(p)
            set_run(p.add_run(f"{i} ـ {stem}:"), size=13.2, bold=True, color=DARK_BLUE)
            cp = doc.add_paragraph()
            set_bidi_paragraph(cp)
            set_run(cp.add_run("     ـ     ".join(choices)), size=12.5, color=BLACK)
            if answer_mode:
                ap = doc.add_paragraph()
                set_bidi_paragraph(ap)
                set_run(ap.add_run("الإجابة: "), size=12.5, bold=True, color=GREEN)
                set_run(ap.add_run(answer), size=12.5)
            else:
                lp = doc.add_paragraph("الإجابة: ........................................................")
                set_bidi_paragraph(lp)
                for r in lp.runs:
                    set_run(r, size=11.5, color=GRAY)
    else:
        for i, (q, a) in enumerate(items, 1):
            add_question(doc, str(i), q, a, answer_mode=answer_mode, lines=2)


def build_whole_lesson_exercise():
    return [
        ("اكتب الفكرة العامة للنص في جملة واحدة.", GENERAL_IDEA),
        ("حدد ثلاث صفات للصديق الوفي مع دليل لكل صفة.", "الأمانة: يأتمنهم على أسراره. المساندة: يسعفك في الشدة. الثبات: في السراء والضراء."),
        ("وازن بين الصداقة الحقيقية والصداقة الزائفة.", "الحقيقية أمانة ونصح ومساندة وإيثار وثبات، والزائفة مصلحة وكلام خادع وانقلاب عند تغير الظروف."),
        ("ما دور المشورة والنقد القويم في الصداقة؟", "يكملان الرأي، ويكشفان الخطأ، ويساعدان على الإصلاح دون مجاملة."),
        ("وضح كيف يحافظ الحوار على مشاعر الصديق.", "بالرفق والرقة واللطف والاستماع وتجنب المراء والخصومة."),
        ("لماذا تعد المواقف أقوى من الكلمات في اختيار الصديق؟", "لأن الكلمات قد تخدع، أما المواقف في الشدة والغيبة فتظهر الثبات والصدق."),
        ("حلل بنية النص من البداية إلى الخاتمة.", "بدأ بتعريف الصديق الوفي، ثم حقوقه وآداب التعامل، ثم الإيثار، ثم التحذير من سوء الاختيار، وختم بأثر الصحبة في المستقبل."),
        ("حدد غرض الكاتب والجمهور المستهدف ووسيلتين للإقناع.", "الغرض الإرشاد والإقناع، والجمهور الشباب، والوسائل التفسير والأمثلة والتعليل والأمر والنهي."),
        ("استخرج من النص تضادين وتقاربين.", "السراء والضراء، ودًا وبغضًا. الرقة واللطف، وبغضًا وموجدة."),
        ("وضح جمال الكلمات المعسولة ونحت ملامح شخصيتك.", "الأولى تصور الكلام الخادع كشيء حلو جذاب، والثانية تصور أثر الصحبة كعمل نحات يشكل الشخصية تدريجيًا."),
        ("ما العاطفة المسيطرة والجو العام؟", "تقدير للوفاء وحرص على الشباب ونفور من الخداع، في جو إرشادي دافئ مع تحذير."),
        ("حدد أسلوبين مختلفين وغرض كل منهما.", "لا تحاول نهي للنصح والتحذير، ودقق وراع أمران للحث والإرشاد، ويا بني نداء للعطف والتنبيه."),
        ("استخرج قيمة وطبقها في موقف مدرسي.", "حفظ الغيبة: أرفض السخرية من زميلي في غيابه وأدافع عنه دون نشر أسراره."),
        ("اكتب حكمًا نقديًا على أفكار النص.", "أفكار واضحة ومتوازنة وقابلة للتطبيق، جمعت بين اختيار الصديق وأداء حقوقه، واعتمدت أدلة واقعية."),
        ("اكتب فقرة من خمسة أسطر بعنوان كن صديقًا وفيًا.", "نموذج: كن أمينًا على السر، صادقًا في النصيحة، حاضرًا وقت الشدة، رفيقًا في الحوار، مخلصًا في العطاء، واحفظ صديقك في غيبته كما تحفظه في حضوره."),
        ("صمم ميثاق صداقة من خمس قواعد.", "حفظ السر، قول الحق بلطف، المساندة، احترام المشاعر، وعدم المن أو استغلال العلاقة."),
        ("طبق معيار النص على صداقة رقمية.", "التحقق من الهوية، عدم مشاركة الخصوصية سريعًا، ملاحظة ثبات الاحترام، ورفض الضغط أو الابتزاز."),
        ("ما أثر الصحبة الصالحة في النجاح؟", "تنظم الوقت وتشجع على العلم والانضباط وتدعم الإنسان وقت الصعوبة وتقوم سلوكه."),
        ("كيف يكون الإنسان الصديق الذي يبحث عنه؟", "بأن يطبق على نفسه الأمانة والوفاء والرفق والإيثار والثبات وحسن النصح."),
        ("اكتب المغزى العام بأسلوبك.", "اختيار الصديق قرار يصنع الشخصية والمستقبل، والصداقة الحقيقية حقوق ومواقف لا كلمات ومصالح.")
    ]


def add_end_exercises(doc, answer_mode):
    add_numbered_exercise(doc, "التدريبات الشاملة في نهاية الدرس ـ اختر الإجابة الصحيحة", MCQ, answer_mode, kind="mcq")
    add_numbered_exercise(doc, "أكمل من ألفاظ النص ـ عشرة أسئلة", COMPLETE_TEXT, answer_mode)
    add_numbered_exercise(doc, "أكمل بالمعنى المناسب ـ عشرة أسئلة", COMPLETE_MEANING, answer_mode)
    add_numbered_exercise(doc, "أجب عن الأسئلة الآتية ـ عشرة أسئلة", GENERAL_QA, answer_mode)
    add_numbered_exercise(doc, "تدريب شامل على عناصر تحليل النص", ANALYSIS_Q, answer_mode)
    add_numbered_exercise(doc, "تدريب شامل على عناصر تذوق النص", TASTE_Q, answer_mode)
    add_numbered_exercise(doc, "أسئلة كتاب المدرسة مجمعة ومدمجة", SCHOOLBOOK_END, answer_mode)

    add_title_band(doc, "تدريبات شاملة مستقلة على كل فقرة", fill=PURPLE, page_break_before=True)
    for idx, pdata in enumerate(PARAGRAPHS, 1):
        add_subtitle(doc, f"التدريب الشامل على الفقرة {idx}", color=PURPLE, size=16)
        for qnumb, (q, a) in enumerate(paragraph_exercise_items(pdata, idx), 1):
            source = "كتاب المدرسة" if qnumb == 4 else None
            add_question(doc, f"{idx}.{qnumb}", q, a, source=source, answer_mode=answer_mode, lines=2)

    add_numbered_exercise(doc, "التدريب النهائي الشامل على الدرس كله", build_whole_lesson_exercise(), answer_mode)


def validate_docx(path: Path, answer_mode: bool):
    assert path.exists() and path.stat().st_size > 50000, f"DOCX missing or too small: {path}"
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
    required = [
        "عناصر تحليل النص مطبقة على الفقرة",
        "عناصر تذوق النص مطبقة على الفقرة",
        "أسئلة كتاب المدرسة المدمجة",
        "تدريبات شاملة مستقلة على كل فقرة",
        "التدريب النهائي الشامل على الدرس كله",
        "١ ـ نوع النص",
        "١٣ ـ الموازنة الشعرية",
        "الكلمة أو التعبير",
        "المعنى",
        "المضاد",
        "الجمع",
        "المفرد",
    ]
    for phrase in required:
        assert phrase in xml, f"Missing required phrase: {phrase}"
    assert xml.count("w:bidiVisual") >= 80, "RTL table markers are insufficient"
    assert xml.count("أسئلة كتاب المدرسة المدمجة") == len(PARAGRAPHS), "Book questions not merged under every paragraph"
    assert xml.count("عناصر تحليل النص مطبقة على الفقرة") == len(PARAGRAPHS), "Analysis not placed after every paragraph"
    assert xml.count("عناصر تذوق النص مطبقة على الفقرة") == len(PARAGRAPHS), "Taste not placed after every paragraph"
    if answer_mode:
        assert xml.count("الإجابة:") >= 150, "Answer key is incomplete"


def build(out_path: Path, answer_mode: bool):
    doc = Document()
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
    normal.font.size = Pt(13)

    add_cover(doc, answer_mode)
    add_intro(doc)
    for idx, pdata in enumerate(PARAGRAPHS, 1):
        add_paragraph_unit(doc, pdata, idx, answer_mode)
    add_end_exercises(doc, answer_mode)
    add_header_footer(doc, answer_mode)

    core = doc.core_properties
    core.title = f"{LESSON_TITLE} ـ {'نموذج الإجابة' if answer_mode else 'نسخة الطالب'}"
    core.subject = "شرح وتحليل وتذوق وتدريبات للصف الثالث الإعدادي"
    core.author = "إعداد تعليمي"
    core.keywords = "الصداقة، الصف الثالث الإعدادي، تحليل النص، تذوق النص، كتاب المدرسة"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    validate_docx(out_path, answer_mode)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--outdir", default="final-lesson3/output")
    args = parser.parse_args()
    outdir = Path(args.outdir)
    student = outdir / "الصداقة_النظام_النهائي_المعتمد_نسخة_الطالب_Word2010.docx"
    answer = outdir / "نموذج_إجابة_الصداقة_النظام_النهائي_المعتمد_Word2010.docx"
    build(student, False)
    build(answer, True)
    print(student)
    print(answer)


if __name__ == "__main__":
    main()
