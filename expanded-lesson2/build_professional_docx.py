#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import argparse
import os
import re
from typing import Iterable, Optional

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT_NAME = "Arial"
BODY_COLOR = "242424"
TEAL = "147C91"
LIGHT_TEAL = "E6F5F8"
LIGHT_MINT = "E6F7F3"
ORANGE = "F4A261"
LIGHT_ORANGE = "FFF6EB"
RED = "D9272E"
LIGHT_RED = "FFF0F0"
GREEN = "2A9D8F"
LIGHT_BLUE = "EDF8FC"
GOLD = "B17D00"
LIGHT_GOLD = "FFF8DA"
BORDER = "9CBFC8"
WHITE = "FFFFFF"


def rgb(hex_color: str) -> RGBColor:
    value = hex_color.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def set_run_font(run, size: float = 14, bold: Optional[bool] = None,
                 color: str = BODY_COLOR, italic: Optional[bool] = None) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(key), FONT_NAME)
    if rpr.find(qn("w:rtl")) is None:
        rpr.append(OxmlElement("w:rtl"))


def set_paragraph_rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT,
                      line_spacing: float = 1.42, space_before: float = 0,
                      space_after: float = 4, keep_with_next: bool = False) -> None:
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.line_spacing = line_spacing
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.keep_with_next = keep_with_next
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        ppr.append(OxmlElement("w:bidi"))


def set_paragraph_shading(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_bottom_border(paragraph, color: str, size: str = "12") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)


def set_cell_shading(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER, size: str = "6") -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, twips: int) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:w"), str(twips))
    tcw.set(qn("w:type"), "dxa")


def set_table_rtl(table) -> None:
    tblpr = table._tbl.tblPr
    if tblpr.find(qn("w:bidiVisual")) is None:
        tblpr.append(OxmlElement("w:bidiVisual"))


def repeat_table_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trpr.append(header)


def prevent_row_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, 10.5, color=RED)


def normalise_text(value: str) -> str:
    value = value.replace("\xa0", " ")
    value = re.sub(r"[\t\r\n]+", " ", value)
    value = re.sub(r" {2,}", " ", value)
    return value


def node_style(node: Tag, base_size: float, base_bold: bool, base_color: str):
    classes = set(node.get("class", []))
    size = 11.5 if "small" in classes or "school" in classes else base_size
    bold = base_bold or node.name in {"b", "strong"} or "bold" in classes or "school" in classes
    italic = node.name in {"i", "em"}
    color = base_color
    if "red" in classes or "school" in classes:
        color = RED
    elif "teal" in classes:
        color = TEAL
    elif "gold" in classes:
        color = GOLD
    return size, bold, color, italic


def add_inline(paragraph, node, size: float = 14, bold: bool = False,
               color: str = BODY_COLOR) -> None:
    if isinstance(node, NavigableString):
        text = normalise_text(str(node))
        if text:
            run = paragraph.add_run(text)
            set_run_font(run, size, bold=bold, color=color)
        return
    if not isinstance(node, Tag):
        return
    if node.name == "br":
        paragraph.add_run().add_break()
        return
    nsize, nbold, ncolor, nitalic = node_style(node, size, bold, color)
    for child in node.children:
        if isinstance(child, NavigableString):
            text = normalise_text(str(child))
            if text:
                run = paragraph.add_run(text)
                set_run_font(run, nsize, bold=nbold, color=ncolor, italic=nitalic)
        elif isinstance(child, Tag):
            add_inline(paragraph, child, nsize, nbold, ncolor)


def add_box(document: Document, tag: Tag, fill: str, border: str,
            size: float = 14, align=WD_ALIGN_PARAGRAPH.RIGHT) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_rtl(table)
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, fill)
    set_cell_border(cell, border, "10")
    set_cell_margins(cell, 110, 140, 110, 140)
    paragraph = cell.paragraphs[0]
    set_paragraph_rtl(paragraph, align, 1.5, 0, 0)
    add_inline(paragraph, tag, size=size)
    spacer = document.add_paragraph()
    set_paragraph_rtl(spacer, space_after=1)


def add_heading(document: Document, text: str, level: int, first_h1: bool = False) -> None:
    text = normalise_text(text).strip()
    if first_h1:
        p = document.add_paragraph()
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.CENTER, 1.15, 2, 6, True)
        run = p.add_run(text)
        set_run_font(run, 17, bold=True, color=TEAL)
        return
    if level == 2:
        p = document.add_paragraph()
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT, 1.2, 8, 8, True)
        set_paragraph_shading(p, LIGHT_TEAL)
        run = p.add_run("  " + text + "  ")
        set_run_font(run, 18, bold=True, color=TEAL)
        return
    if level == 3:
        p = document.add_paragraph()
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT, 1.2, 8, 6, True)
        set_paragraph_shading(p, LIGHT_BLUE)
        set_paragraph_bottom_border(p, TEAL, "8")
        run = p.add_run("  " + text + "  ")
        set_run_font(run, 16, bold=True, color=TEAL)
        return

    lower = text
    if "القاموس" in lower:
        color, fill = ORANGE, LIGHT_ORANGE
    elif "التعبيرات" in lower or "الدلالات" in lower:
        color, fill = RED, LIGHT_RED
    elif "القيم" in lower or "المبادئ" in lower:
        color, fill = GREEN, LIGHT_MINT
    elif "أسئلة" in lower or "تدريبات" in lower:
        color, fill = TEAL, LIGHT_BLUE
    else:
        color, fill = TEAL, WHITE
    p = document.add_paragraph()
    set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT, 1.2, 7, 5, True)
    if fill != WHITE:
        set_paragraph_shading(p, fill)
    set_paragraph_bottom_border(p, color, "9")
    run = p.add_run("  " + text + "  ")
    set_run_font(run, 14.5, bold=True, color=color)


def choose_table_header(rows: list[list[Tag]]) -> str:
    header_text = " ".join(normalise_text(cell.get_text(" ", strip=True)) for cell in rows[0]) if rows else ""
    if "الكلمة" in header_text or "المعنى" in header_text and "المضاد" in header_text:
        return ORANGE
    if "التعبير" in header_text and "الدلالة" in header_text:
        return RED
    if "القيمة" in header_text or "المبدأ" in header_text:
        return GREEN
    return TEAL


def column_ratios(column_count: int, header_text: str) -> list[float]:
    if column_count == 5:
        return [0.18, 0.29, 0.18, 0.18, 0.17]
    if column_count == 3:
        return [0.27, 0.34, 0.39]
    if column_count == 2:
        if "العنصر" in header_text:
            return [0.24, 0.76]
        return [0.38, 0.62]
    return [1.0 / column_count] * column_count


def add_html_table(document: Document, tag: Tag) -> None:
    tr_tags = tag.find_all("tr")
    rows = []
    for tr in tr_tags:
        cells = tr.find_all(["th", "td"], recursive=False)
        if cells:
            rows.append(cells)
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_rtl(table)
    header_color = choose_table_header(rows)
    header_text = " ".join(normalise_text(c.get_text(" ", strip=True)) for c in rows[0])
    ratios = column_ratios(cols, header_text)
    total_twips = 10500

    for r_index, html_cells in enumerate(rows):
        row = table.rows[r_index]
        prevent_row_split(row)
        if r_index == 0:
            repeat_table_header(row)
        for c_index in range(cols):
            cell = row.cells[c_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            set_cell_width(cell, int(total_twips * ratios[c_index]))
            cell.text = ""
            paragraph = cell.paragraphs[0]
            is_header = r_index == 0
            if is_header:
                set_cell_shading(cell, header_color)
                set_paragraph_rtl(paragraph, WD_ALIGN_PARAGRAPH.CENTER, 1.18, 0, 0)
            else:
                set_cell_shading(cell, LIGHT_BLUE if r_index % 2 == 1 else WHITE)
                if cols == 5 and c_index != 1:
                    align = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    align = WD_ALIGN_PARAGRAPH.RIGHT
                set_paragraph_rtl(paragraph, align, 1.28, 0, 0)
            if c_index < len(html_cells):
                add_inline(paragraph, html_cells[c_index], size=12.2,
                           bold=is_header, color=WHITE if is_header else BODY_COLOR)
    spacer = document.add_paragraph()
    set_paragraph_rtl(spacer, space_after=1)


def add_list(document: Document, tag: Tag) -> None:
    ordered = tag.name == "ol"
    for index, li in enumerate(tag.find_all("li", recursive=False), start=1):
        p = document.add_paragraph()
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT, 1.42, 0, 3)
        prefix = f"{index} ـ " if ordered else "• "
        run = p.add_run(prefix)
        set_run_font(run, 13.5, bold=True, color=TEAL)
        add_inline(p, li, size=13.5)


def add_regular_paragraph(document: Document, tag: Tag) -> None:
    classes = set(tag.get("class", []))
    text = normalise_text(tag.get_text(" ", strip=True))
    if not text:
        return
    if "cover-sub" in classes:
        p = document.add_paragraph()
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.CENTER, 1.15, 4, 7, True)
        run = p.add_run(text)
        set_run_font(run, 26, bold=True, color=TEAL)
        return
    if "box" in classes:
        add_box(document, tag, LIGHT_TEAL, TEAL, 13.5,
                WD_ALIGN_PARAGRAPH.CENTER if "center" in classes else WD_ALIGN_PARAGRAPH.RIGHT)
        return
    if "goldbox" in classes:
        add_box(document, tag, LIGHT_GOLD, ORANGE, 13.5)
        return
    if "textbox" in classes:
        add_box(document, tag, "FFFCF7", ORANGE, 14.2)
        return
    if "note" in classes:
        add_box(document, tag, LIGHT_RED, RED, 13.2)
        return

    align = WD_ALIGN_PARAGRAPH.CENTER if "center" in classes else WD_ALIGN_PARAGRAPH.RIGHT
    p = document.add_paragraph()
    set_paragraph_rtl(p, align, 1.45, 0, 4)
    if "q" in classes or text.startswith("س:") or text.startswith("س ـ"):
        add_inline(p, tag, size=13.4, bold=True, color=TEAL)
    elif text.startswith("ج:") or text.startswith("الإجابة"):
        add_inline(p, tag, size=13.2, color=BODY_COLOR)
    elif "lines" in classes:
        run = p.add_run(text)
        set_run_font(run, 12.5, color="777777")
    else:
        add_inline(p, tag, size=13.5)


def configure_document(document: Document, header_title: str, is_answer: bool) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)
    section.header_distance = Cm(0.45)
    section.footer_distance = Cm(0.55)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(13.5)
    rpr = normal._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(key), FONT_NAME)

    header = section.header
    hp = header.paragraphs[0]
    set_paragraph_rtl(hp, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 0, 2)
    hr = hp.add_run(header_title)
    set_run_font(hr, 11.5, bold=True, color=TEAL)
    set_paragraph_bottom_border(hp, TEAL, "8")

    first_header = section.first_page_header
    fhp = first_header.paragraphs[0]
    set_paragraph_rtl(fhp, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 0, 0)

    footer = section.footer
    fp = footer.paragraphs[0]
    set_paragraph_rtl(fp, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 0, 0)
    label = fp.add_run("صفحة ")
    set_run_font(label, 10.5, color=BODY_COLOR)
    add_page_field(fp)

    first_footer = section.first_page_footer
    ffp = first_footer.paragraphs[0]
    set_paragraph_rtl(ffp, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 0, 0)

    document.core_properties.title = header_title
    document.core_properties.subject = "شرح وتدريبات لغة عربية للصف الثالث الإعدادي"
    document.core_properties.author = "OpenAI"
    document.core_properties.keywords = "لغة عربية، أغلى من الذهب، Word 2010"


def build(input_path: str, output_path: str, header_title: str, is_answer: bool = False) -> None:
    with open(input_path, "r", encoding="utf-8") as stream:
        soup = BeautifulSoup(stream.read(), "lxml")

    document = Document()
    if document.paragraphs:
        paragraph = document.paragraphs[0]._element
        paragraph.getparent().remove(paragraph)
    configure_document(document, header_title, is_answer)

    root = soup.find("div", class_="Section1") or soup.body
    first_h1 = True
    for child in root.children:
        if isinstance(child, NavigableString):
            continue
        if not isinstance(child, Tag):
            continue
        classes = set(child.get("class", []))
        if "pagebreak" in classes:
            document.add_page_break()
            continue
        if child.name == "h1":
            add_heading(document, child.get_text(" ", strip=True), 1, first_h1=first_h1)
            first_h1 = False
        elif child.name == "h2":
            add_heading(document, child.get_text(" ", strip=True), 2)
        elif child.name == "h3":
            add_heading(document, child.get_text(" ", strip=True), 3)
        elif child.name == "h4":
            add_heading(document, child.get_text(" ", strip=True), 4)
        elif child.name == "table":
            add_html_table(document, child)
        elif child.name in {"ul", "ol"}:
            add_list(document, child)
        elif child.name in {"p", "div"}:
            add_regular_paragraph(document, child)

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    document.save(output_path)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build professionally formatted Arabic DOCX from Word-compatible HTML")
    parser.add_argument("input")
    parser.add_argument("output")
    parser.add_argument("--title", required=True)
    parser.add_argument("--answer", action="store_true")
    args = parser.parse_args()
    build(args.input, args.output, args.title, args.answer)
